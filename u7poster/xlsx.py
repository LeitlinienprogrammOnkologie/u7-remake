"""assets.xlsx writer (openpyxl). All coordinate conversions are live formulas
that reference the Config sheet."""
from __future__ import annotations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .config import CONFIG, CONFIG_NOTES

STATUS_LIST = "open,generating,done,skipped"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
FORMULA_FILL = PatternFill("solid", fgColor="F2F2F2")


def _config_sheet(wb: Workbook, cfg: dict):
    ws = wb.active
    ws.title = "Config"
    ws.append(["key", "value", "note"])
    addr = {}
    for i, (k, v) in enumerate(cfg.items(), start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        ws.cell(i, 3, CONFIG_NOTES.get(k, ""))
        addr[k] = f"Config!$B${i}"
        wb.defined_names[k] = DefinedName(k, attr_text=f"Config!$B${i}")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws.freeze_panes = "A2"
    return addr


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 2)}"


def _widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _coord_formulas(A, r, xcol, ycol):
    """frac, base px, up px, coarse px formulas for a row r given tile x/y column letters."""
    fx = f"={xcol}{r}/{A['map_tiles']}"
    fy = f"={ycol}{r}/{A['map_tiles']}"
    return [fx, fy,
            f"=ROUND({xcol}{r}/{A['map_tiles']}*{A['base_px']},1)", f"=ROUND({ycol}{r}/{A['map_tiles']}*{A['base_px']},1)",
            f"=ROUND({xcol}{r}/{A['map_tiles']}*{A['up_px']},1)", f"=ROUND({ycol}{r}/{A['map_tiles']}*{A['up_px']},1)",
            f"=ROUND({A['coarse_offset_x']}+{xcol}{r}*{A['coarse_scale_x']},1)",
            f"=ROUND({A['coarse_offset_y']}+{ycol}{r}*{A['coarse_scale_y']},1)"]


COORD_HEADERS = ["frac_x", "frac_y", "base_px_x", "base_px_y", "up_px_x", "up_px_y", "coarse_px_x", "coarse_px_y"]


def build_workbook(path: Path, cfg: dict, landmarks, npc_rows, monster_rows, defects, crop_paths: dict) -> dict:
    """Writes the workbook; returns {sheet: {id: suggested_cm_value}} computed in Python for the overlays."""
    wb = Workbook()
    A = _config_sheet(wb, cfg)
    status_dv = DataValidation(type="list", formula1=f'"{STATUS_LIST}"', allow_blank=True)
    prio_dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    cm_out = {"landmarks": {}, "npcs": {}, "monsters": {}}

    # ---------------- Landmarks ----------------
    ws = wb.create_sheet("Landmarks")
    head = ["id", "name", "category", "tile_x", "tile_y"] + COORD_HEADERS + \
           ["footprint_tiles_w", "footprint_tiles_h", "suggested_print_cm", "priority", "prompt_hint", "status",
            "name_source", "confidence", "npcs_here", "signs_here"]
    ws.append(head)
    ws.add_data_validation(status_dv)
    ws.add_data_validation(prio_dv)
    for i, l in enumerate(landmarks, start=2):
        r = i
        row = [l.id, l.name, l.category, l.tile_x, l.tile_y] + _coord_formulas(A, r, "D", "E") + [l.w, l.h]
        cm_formula = (f"=ROUND(MAX(CHOOSE(Q{r},{A['landmark_min_cm_p1']},{A['landmark_min_cm_p2']},{A['landmark_min_cm_p3']}),"
                      f"MAX(N{r},O{r})*{A['poster_w_cm']}/{A['map_tiles']}*{A['landmark_footprint_factor']}),1)")
        row += [cm_formula, l.priority, l.prompt_hint, "open", l.name_source, l.confidence,
                ", ".join(l.npcs)[:250], " | ".join(l.signs)[:250]]
        ws.append(row)
        pmin = cfg[f"landmark_min_cm_p{l.priority}"]
        cm_out["landmarks"][l.id] = round(max(pmin, max(l.w, l.h) * cfg["poster_w_cm"] / cfg["map_tiles"] * cfg["landmark_footprint_factor"]), 1)
        status_dv.add(f"R{r}")
        prio_dv.add(f"Q{r}")
        for c in range(6, 14):
            ws.cell(r, c).fill = FORMULA_FILL
        ws.cell(r, 16).fill = FORMULA_FILL
    _style_header(ws, len(head))
    _widths(ws, {"A": 7, "B": 34, "C": 11, "P": 12, "R": 46, "S": 10, "T": 28, "U": 10, "V": 40, "W": 40})

    # ---------------- NPCs ----------------
    ws = wb.create_sheet("NPCs")
    head = ["id", "name", "shape", "frame", "tile_x", "tile_y", "home_source"] + COORD_HEADERS + \
           ["role", "suggested_print_cm", "priority", "appearance_notes", "status", "sex", "start_tile_x", "start_tile_y",
            "nearest_landmark_id", "nearest_landmark"]
    ws.append(head)
    sdv = DataValidation(type="list", formula1=f'"{STATUS_LIST}"', allow_blank=True)
    pdv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(sdv)
    ws.add_data_validation(pdv)
    for i, n in enumerate(npc_rows, start=2):
        r = i
        row = [n.id, n.name, n.shape, n.frame, n.home_x, n.home_y, n.home_source] + _coord_formulas(A, r, "E", "F")
        row += [n.role, f"=CHOOSE(R{r},{A['npc_cm_p1']},{A['npc_cm_p2']},{A['npc_cm_p3']})", n.priority, "", "open",
                "female" if n.female else "male", n.start_x, n.start_y, n.landmark_id, n.landmark_name]
        ws.append(row)
        cm_out["npcs"][n.id] = cfg[f"npc_cm_p{n.priority}"]
        sdv.add(f"T{r}")
        pdv.add(f"R{r}")
        for c in range(8, 16):
            ws.cell(r, c).fill = FORMULA_FILL
        ws.cell(r, 17).fill = FORMULA_FILL
    _style_header(ws, len(head))
    _widths(ws, {"A": 7, "B": 16, "G": 22, "P": 40, "S": 40, "T": 10, "Y": 30})

    # ---------------- Monsters ----------------
    ws = wb.create_sheet("Monsters")
    head = ["id", "shape", "shape_name", "kind", "count_eggs", "spawn_max", "tile_x", "tile_y"] + COORD_HEADERS + \
           ["region_id", "region_label", "suggested_print_cm", "priority", "status", "alignment", "spread_tiles_w",
            "spread_tiles_h", "under_mountains"]
    ws.append(head)
    sdv = DataValidation(type="list", formula1=f'"{STATUS_LIST}"', allow_blank=True)
    pdv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(sdv)
    ws.add_data_validation(pdv)
    for i, m in enumerate(monster_rows, start=2):
        r = i
        row = [m.id, m.shape, m.shape_name, m.kind, m.count_eggs, m.spawn_max, m.tile_x, m.tile_y] + _coord_formulas(A, r, "G", "H")
        row += [m.region_id, m.region_name, f"=CHOOSE(T{r},{A['monster_cm_p1']},{A['monster_cm_p2']},{A['monster_cm_p3']})",
                m.priority, "open", m.align, m.bbox_w, m.bbox_h, "yes" if m.in_dungeon else ""]
        ws.append(row)
        cm_out["monsters"][m.id] = cfg[f"monster_cm_p{m.priority}"]
        sdv.add(f"U{r}")
        pdv.add(f"T{r}")
        for c in range(9, 17):
            ws.cell(r, c).fill = FORMULA_FILL
        ws.cell(r, 19).fill = FORMULA_FILL
    _style_header(ws, len(head))
    _widths(ws, {"A": 7, "C": 18, "R": 30, "U": 10})

    # ---------------- FixLog ----------------
    ws = wb.create_sheet("FixLog")
    head = ["id", "tile_row", "tile_col", "centre_frac_x", "centre_frac_y", "box_px", "cause", "issue", "fix_type", "prompt",
            "tool", "status", "before_crop", "after_crop", "class", "centre_px_x", "centre_px_y", "box_w_px", "box_h_px",
            "area_px", "tile_x", "tile_y"]
    ws.append(head)
    dvs = {
        "cause": DataValidation(type="list", formula1='"base,upscale"', allow_blank=True),
        "fix_type": DataValidation(type="list", formula1='"regenerate,paint,re-upscale tile"', allow_blank=True),
        "tool": DataValidation(type="list", formula1='"Photoshop,Midjourney,Magnific"', allow_blank=True),
        "status": DataValidation(type="list", formula1='"open,done,skipped"', allow_blank=True),
    }
    for dv in dvs.values():
        ws.add_data_validation(dv)
    for i, d in enumerate(defects, start=2):
        r = i
        fix = "paint" if d.dclass in ("water-land", "river", "road", "structure") else "regenerate"
        row = [d.id,
               f"=INT(E{r}*{A['fix_grid_rows']})+1", f"=INT(D{r}*{A['fix_grid_cols']})+1",
               f"=ROUND(P{r}/{A['coarse_px']},4)", f"=ROUND(Q{r}/{A['coarse_px']},4)",
               f"=MAX(R{r},S{r})", "upscale", d.issue, fix, d.prompt, "", "", crop_paths.get(d.id, ""), "",
               d.dclass, d.cx, d.cy, d.w, d.h, d.area, d.tile_x, d.tile_y]
        ws.append(row)
        for c in (2, 3, 4, 5, 6):
            ws.cell(r, c).fill = FORMULA_FILL
    last = max(len(defects) + 1, 2) + 500   # leave validated empty rows for manual entries
    dvs["cause"].add(f"G2:G{last}")
    dvs["fix_type"].add(f"I2:I{last}")
    dvs["tool"].add(f"K2:K{last}")
    dvs["status"].add(f"L2:L{last}")
    _style_header(ws, len(head))
    _widths(ws, {"A": 7, "H": 60, "J": 55, "K": 12, "L": 9, "M": 18, "N": 18})

    wb.save(path)
    return cm_out
